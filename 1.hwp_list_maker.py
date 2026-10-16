# -*- coding: utf-8 -*-
r"""
인쇄 목록 자동 생성 스크립트 (1.hwp_list_maker.py)
==================================================

루트 폴더를 지정하면 그 안의 모든 하위 폴더를 탐색해 .hwp 파일을 찾고,
2.hwp_print.py가 읽는 엑셀 목록 형식으로 만들어줍니다.

    1차경로 = 지정한 루트 폴더
    2차경로 = 루트에서 파일까지의 하위 폴더 경로 (예: C01_01_2_창녕군청\2. 수도과)
    파일명  = hwp 파일 이름
    부수    = 기본값 1 (생성 후 엑셀에서 필요한 값으로 수정)

사용법:
    python 1.hwp_list_maker.py "C:\경로\2026년 8월 청구리스트"

저장 위치와 파일명:
    - 저장 폴더는 아래 OUTPUT_DIR 설정값을 따릅니다.
      (기본값: 이 스크립트가 있는 폴더 안의 print_list 폴더)
    - 파일명은 "print_list_루트폴더명.xlsx" 형식으로 자동 결정됩니다.
      예) 루트가 ...\2026년 8월 청구리스트 이면
          → print_list_2026년 8월 청구리스트.xlsx

옵션:
    --output "경로.xlsx"   저장 위치/이름을 직접 지정 (OUTPUT_DIR 무시)
    --copies N             부수 기본값 (기본 1)

주의:
    - 같은 이름의 목록 파일이 이미 있으면 덮어쓰지 않고
      뒤에 날짜시간을 붙인 새 이름으로 저장합니다.
      (이미 부수를 수정해 둔 목록을 실수로 날리지 않기 위함)
    - 엑셀 임시 파일(~$로 시작)은 무시합니다.

전형적인 회차 워크플로우:
    1) 이전 회차 폴더 복사 → 새 회차 폴더 만들기
    2) 문서 내용 수정
    3) 이 스크립트로 새 회차 폴더의 인쇄 목록 생성
    4) 엑셀에서 부수만 수정 (인쇄 안 할 파일은 0 입력) 후 저장
    5) python 2.hwp_print.py "목록.xlsx" --dry-run 으로 확인 후 인쇄
"""

import argparse
import datetime
import os
from pathlib import Path

# ============================================================
# 생성된 엑셀 목록이 저장될 폴더
#   - 기본값: 이 스크립트가 있는 폴더 안의 print_list 폴더
#     (어느 PC, 어느 위치로 옮겨도 그대로 동작)
#   - 다른 곳에 저장하려면 예: OUTPUT_DIR = r"C:\Users\사용자명\Documents\print_list"
#   - 빈 문자열("")로 두면 루트 폴더 안에 저장됩니다.
# ============================================================
OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "print_list")


def find_hwp_files(root):
    """루트 폴더 아래 모든 .hwp 파일을 (하위폴더 상대경로, 파일명) 목록으로 반환"""
    entries = []
    for path in sorted(root.rglob("*.hwp")):
        if not path.is_file():
            continue
        if path.name.startswith("~$"):
            continue  # 오피스/한글 임시 잠금 파일 무시
        rel_dir = path.parent.relative_to(root)
        sub_path = os.sep.join(rel_dir.parts)  # 루트 바로 아래면 빈 문자열
        entries.append((sub_path, path.name))
    return entries


def unique_output_path(output):
    """이미 존재하는 파일을 덮어쓰지 않도록 새 이름을 만들어 반환"""
    if not output.exists():
        return output
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    return output.with_name(f"{output.stem}_{ts}{output.suffix}")


def write_xlsx(entries, root, output, default_copies):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise SystemExit(
            "엑셀 파일 생성에는 openpyxl이 필요합니다.\n"
            "명령 프롬프트에서 'pip install openpyxl' 을 실행한 뒤 다시 시도해 주세요."
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "인쇄목록"

    ws.append(["1차경로", "2차경로", "파일명", "부수"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")

    root_str = str(root)
    for sub_path, filename in entries:
        ws.append([root_str, sub_path, filename, default_copies])

    for i, width in enumerate([70, 35, 32, 8], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    wb.save(output)


def main():
    parser = argparse.ArgumentParser(
        description="루트 폴더의 모든 하위 폴더에서 hwp 파일을 찾아 인쇄 목록 엑셀을 생성합니다."
    )
    parser.add_argument("root", help="hwp 파일들을 탐색할 루트 폴더")
    parser.add_argument("--output", help="저장할 엑셀 파일 경로 (생략 시 OUTPUT_DIR 폴더에 print_list_루트폴더명.xlsx)")
    parser.add_argument("--copies", type=int, default=1, help="부수 기본값 (기본 1)")
    args = parser.parse_args()

    root = Path(args.root).expanduser().resolve()
    if not root.is_dir():
        raise SystemExit(f"폴더를 찾을 수 없습니다: {root}")

    entries = find_hwp_files(root)
    if not entries:
        raise SystemExit(f"'{root}' 아래에서 .hwp 파일을 찾지 못했습니다.")

    if args.output:
        output = Path(args.output).expanduser().resolve()
        if output.suffix.lower() != ".xlsx":
            output = output.with_suffix(".xlsx")
    else:
        # 저장 폴더: OUTPUT_DIR 설정값 (비어있으면 루트 폴더)
        if OUTPUT_DIR.strip():
            out_dir = Path(OUTPUT_DIR).expanduser()
            try:
                out_dir.mkdir(parents=True, exist_ok=True)
            except Exception as e:
                raise SystemExit(
                    f"저장 폴더를 만들거나 접근할 수 없습니다: {out_dir}\n({e})\n"
                    "코드 상단의 OUTPUT_DIR 경로를 확인해 주세요."
                )
        else:
            out_dir = root
        # 파일명: print_list_루트폴더명.xlsx
        output = out_dir / f"print_list_{root.name}.xlsx"
    output = unique_output_path(output)

    write_xlsx(entries, root, output, args.copies)

    print(f"\n루트 폴더: {root}")
    print(f"찾은 hwp 파일: {len(entries)}개")
    print("-" * 60)
    for sub_path, filename in entries:
        shown = f"{sub_path}{os.sep}{filename}" if sub_path else filename
        print(f"  {shown}")
    print("-" * 60)
    print(f"목록 저장 완료: {output}")
    print(f"\n다음 단계:")
    print(f"  1. 엑셀에서 위 파일을 열어 '부수' 열을 수정하고 저장하세요 (인쇄 제외할 파일은 0).")
    print(f"  2. python 2.hwp_print.py \"{output}\" --dry-run 으로 확인 후 인쇄하세요.")


if __name__ == "__main__":
    main()